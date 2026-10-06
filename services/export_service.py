from pathlib import Path

from services.timetable_service import DAYS, timetable_rows
from services.timetable_view_service import DEFAULT_FACULTY_COLOUR, campus_timetable_grid
from services.faculty_style_service import get_faculty_style

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except Exception:  # pragma: no cover
    Workbook = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, Image
except Exception:  # pragma: no cover
    SimpleDocTemplate = None


def _version_meta(conn, version_id):
    row = conn.execute(
        """SELECT tv.version_number,tv.status,s.academic_year,'Semester ' || s.semester_number AS semester_name
           FROM timetable_versions tv JOIN semesters s ON s.semester_id=tv.semester_id
           WHERE tv.version_id=?""",
        (version_id,),
    ).fetchone()
    return dict(row) if row else {"version_number": version_id, "status": "", "academic_year": "", "semester_name": ""}


def _grid(conn, version_id):
    return campus_timetable_grid(conn, timetable_rows(conn, version_id))


def _safe_hex(value):
    value = (value or DEFAULT_FACULTY_COLOUR).strip()
    if not value:
        value = DEFAULT_FACULTY_COLOUR
    if not value.startswith("#"):
        value = "#" + value
    if len(value) != 7:
        return DEFAULT_FACULTY_COLOUR
    return value.upper()


def _entry_colour(entry):
    if not entry:
        return "#FFFFFF"
    code = (entry.get("faculty_code") or "").strip().upper()
    if code:
        return get_faculty_style(code)["background"]
    return _safe_hex(entry.get("colour_hex"))


def _cell_text(entries):
    if not entries:
        return ""
    seen = []
    for item in entries:
        code = item.get("module_code") or ""
        if code and code not in seen:
            seen.append(code)
    return " / ".join(seen)


def export_excel(conn, version_id, output_path):
    if Workbook is None:
        raise RuntimeError("openpyxl is required for Excel export.")
    grid = _grid(conn, version_id)
    meta = _version_meta(conn, version_id)
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="C96F16")
    day_fill = PatternFill("solid", fgColor="FFF200")
    neutral_fill = PatternFill("solid", fgColor=DEFAULT_FACULTY_COLOUR.replace("#", ""))

    for day_section in grid["days"]:
        ws = wb.create_sheet(day_section["day"])
        ws.merge_cells("A1:L1")
        ws["A1"] = f"Sol Plaatje University — {meta['academic_year']} {meta['semester_name']} Timetable · {meta['status']}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws["A1"].fill = day_fill
        ws.merge_cells("A2:L2")
        ws["A2"] = day_section["day"].upper()
        ws["A2"].font = Font(bold=True)
        ws["A2"].alignment = Alignment(horizontal="center")
        ws["A2"].fill = day_fill

        headers = ["VENUE", "CAPACITY"] + [f"P{p['period_number']}" for p in grid["periods"]]
        times = ["", ""] + [f"{p['start_time']}-{p['end_time']}" for p in grid["periods"]]
        for c, value in enumerate(headers, 1):
            cell = ws.cell(4, c, value)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = header_fill
            cell.border = border
        for c, value in enumerate(times, 1):
            cell = ws.cell(5, c, value)
            cell.font = Font(bold=True, size=8)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = header_fill
            cell.border = border

        row_index = 6
        for venue in day_section["venues"]:
            venue_name = venue.get("room_name") or venue.get("room_code") or ""
            ws.cell(row_index, 1, venue_name)
            ws.cell(row_index, 2, venue.get("capacity"))
            for c in (1, 2):
                ws.cell(row_index, c).alignment = Alignment(vertical="center", wrap_text=True)
                ws.cell(row_index, c).border = border
            for period in range(1, 11):
                entries = venue["cells"][period]
                cell = ws.cell(row_index, period + 2, _cell_text(entries))
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
                colour = _entry_colour(entries[0]) if entries else "#FFFFFF"
                cell.fill = PatternFill("solid", fgColor=colour.replace("#", "")) if entries else PatternFill("solid", fgColor="FFFFFF")
            row_index += 1

        ws.freeze_panes = "C6"
        ws.column_dimensions["A"].width = 27
        ws.column_dimensions["B"].width = 11
        for col in range(3, 13):
            ws.column_dimensions[get_column_letter(col)].width = 18
        for row in range(6, row_index):
            ws.row_dimensions[row].height = 30
        ws.sheet_view.showGridLines = False

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def export_pdf(conn, version_id, output_path):
    if SimpleDocTemplate is None:
        raise RuntimeError("reportlab is required for PDF export.")
    grid = _grid(conn, version_id)
    meta = _version_meta(conn, version_id)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(output_path), pagesize=landscape(A4),
        rightMargin=12, leftMargin=12, topMargin=14, bottomMargin=14,
    )
    styles = getSampleStyleSheet()
    compact = ParagraphStyle("CompactCell", parent=styles["BodyText"], fontSize=5.4, leading=6.2, alignment=TA_CENTER)
    story = []
    logo = Path(__file__).resolve().parent.parent / "static" / "images" / "spu-logo.png"
    if logo.exists():
        try:
            story.append(Image(str(logo), width=48, height=48))
        except Exception:
            pass
    story.append(Paragraph("Sol Plaatje University", styles["Title"]))
    story.append(Paragraph(
        f"{meta['academic_year']} {meta['semester_name']} Timetable — Version {meta['version_number']} · {meta['status']}",
        styles["Heading2"],
    ))
    story.append(Spacer(1, 6))

    col_widths = [75, 38] + [65] * 10
    for day_index, day_section in enumerate(grid["days"]):
        if day_index:
            story.append(PageBreak())
        story.append(Paragraph(day_section["day"].upper(), styles["Heading2"]))
        header = ["VENUE", "CAPACITY"] + [f"P{p['period_number']}" for p in grid["periods"]]
        times = ["", ""] + [f"{p['start_time']}-{p['end_time']}" for p in grid["periods"]]
        data = [header, times]
        fills = []
        for venue in day_section["venues"]:
            row = [Paragraph(str(venue.get("room_name") or venue.get("room_code") or ""), compact), str(venue.get("capacity") or "")]
            row_fills = []
            for period in range(1, 11):
                entries = venue["cells"][period]
                row.append(Paragraph(_cell_text(entries), compact))
                row_fills.append(_entry_colour(entries[0]) if entries else "#FFFFFF")
            data.append(row)
            fills.append(row_fills)
        table = Table(data, colWidths=col_widths, repeatRows=2, hAlign="CENTER")
        commands = [
            ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#C96F16")),
            ("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 1), 6),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.black),
            ("FONTSIZE", (0, 2), (-1, -1), 5.2),
        ]
        for r_idx, row_fills in enumerate(fills, start=2):
            for p_idx, colour in enumerate(row_fills, start=2):
                commands.append(("BACKGROUND", (p_idx, r_idx), (p_idx, r_idx), colors.HexColor(colour)))
        table.setStyle(TableStyle(commands))
        story.append(table)
    doc.build(story)
    return output_path


def _published_for_user(conn, version_id, user):
    if user and user.get("role") != "ADMIN":
        row = conn.execute("SELECT status FROM timetable_versions WHERE version_id=?", (version_id,)).fetchone()
        if row is None or row["status"] != "PUBLISHED":
            raise ValueError("Only a published timetable can be exported.")


def export_personal_excel(conn, version_id, user, output_path):
    if Workbook is None:
        raise RuntimeError("openpyxl is required for Excel export.")
    _published_for_user(conn, version_id, user)
    rows = timetable_rows(conn, version_id, user)
    wb = Workbook()
    ws = wb.active
    ws.title = "My Timetable"
    ws.merge_cells("A1:F1")
    ws["A1"] = "Sol Plaatje University — My Timetable"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["Period"] + DAYS)
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    by_key = {(day, p): [] for day in DAYS for p in range(1, 11)}
    for r in rows:
        duration = max(1, int(r.get("duration_slots") or 1))
        for p in range(int(r["period_number"]), min(10, int(r["period_number"]) + duration - 1) + 1):
            by_key[(r["day_of_week"], p)].append(r)
    for p in range(1, 11):
        row_values = [f"P{p}"]
        for day in DAYS:
            items = by_key[(day, p)]
            row_values.append("\n".join(dict.fromkeys(i["module_code"] for i in items)))
        ws.append(row_values)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions["A"].width = 12
    for col in range(2, 7):
        ws.column_dimensions[get_column_letter(col)].width = 24
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def export_personal_pdf(conn, version_id, user, output_path):
    if SimpleDocTemplate is None:
        raise RuntimeError("reportlab is required for PDF export.")
    _published_for_user(conn, version_id, user)
    rows = timetable_rows(conn, version_id, user)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(output_path), pagesize=landscape(A4), rightMargin=18, leftMargin=18, topMargin=18, bottomMargin=18)
    styles = getSampleStyleSheet()
    story = [Paragraph("Sol Plaatje University — My Timetable", styles["Title"]), Spacer(1, 8)]
    data = [["Period"] + DAYS]
    by_key = {(day, p): [] for day in DAYS for p in range(1, 11)}
    for r in rows:
        duration = max(1, int(r.get("duration_slots") or 1))
        for p in range(int(r["period_number"]), min(10, int(r["period_number"]) + duration - 1) + 1):
            by_key[(r["day_of_week"], p)].append(r)
    compact = ParagraphStyle("PersonalCell", parent=styles["BodyText"], fontSize=7, leading=8, alignment=TA_CENTER)
    for p in range(1, 11):
        row = [f"P{p}"]
        for day in DAYS:
            items = by_key[(day, p)]
            text = "<br/>".join(dict.fromkeys(i["module_code"] for i in items))
            row.append(Paragraph(text, compact))
        data.append(row)
    table = Table(data, repeatRows=1, colWidths=[45] + [145] * 5)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C96F16")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))
    story.append(table)
    doc.build(story)
    return output_path
