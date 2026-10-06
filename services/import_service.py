import csv
from io import StringIO, BytesIO
from pathlib import Path

from services.setup_service import create_offering

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


DATASETS = {
    "rooms": {
        "required": ["room_code", "room_name", "building_code", "capacity", "room_type"],
    },
    "staff": {
        "required": ["staff_number", "first_name", "last_name", "department_code", "email"],
    },
    "modules": {
        "required": ["module_code", "module_name", "department_code", "credits"],
    },
    "student_groups": {
        "required": ["programme_code", "year_level", "number_of_students"],
    },
    "offerings": {
        "required": ["module_code", "expected_students", "groups", "staff_numbers", "session_type", "sessions_per_week", "duration_slots", "required_room_type", "can_split_duration"],
    },
}


def _read_rows(filename, raw_bytes):
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = raw_bytes.decode("utf-8-sig")
        return list(csv.DictReader(StringIO(text)))
    if suffix in {".xlsx", ".xlsm"}:
        if load_workbook is None:
            raise RuntimeError("openpyxl is required for Excel imports.")
        wb = load_workbook(BytesIO(raw_bytes), read_only=True, data_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(v).strip() if v is not None else "" for v in values[0]]
        rows = []
        for values_row in values[1:]:
            rows.append({headers[i]: values_row[i] for i in range(min(len(headers), len(values_row)))})
        return rows
    raise ValueError("Only .csv and .xlsx files are supported.")


def preview_import(conn, dataset, filename, raw_bytes):
    if dataset not in DATASETS:
        raise ValueError("Unsupported dataset type.")
    rows = _read_rows(filename, raw_bytes)
    required = DATASETS[dataset]["required"]
    errors = []
    warnings = []
    normalized = []

    for index, original in enumerate(rows, start=2):
        row = {str(k).strip(): ("" if v is None else str(v).strip()) for k, v in original.items()}
        missing = [field for field in required if not row.get(field)]
        if missing:
            errors.append(f"Row {index}: missing {', '.join(missing)}")
            continue

        if dataset == "rooms":
            try:
                if int(row["capacity"]) <= 0:
                    raise ValueError
            except ValueError:
                errors.append(f"Row {index}: capacity must be a positive integer.")
                continue
            if conn.execute("SELECT 1 FROM rooms WHERE room_code=?", (row["room_code"],)).fetchone():
                errors.append(f"Row {index}: duplicate room code {row['room_code']}.")
            if not conn.execute("SELECT 1 FROM buildings WHERE building_code=?", (row["building_code"],)).fetchone():
                errors.append(f"Row {index}: unknown building {row['building_code']}.")

        elif dataset == "staff":
            if conn.execute("SELECT 1 FROM staff WHERE staff_number=?", (row["staff_number"],)).fetchone():
                errors.append(f"Row {index}: duplicate staff number {row['staff_number']}.")
            matches = conn.execute("SELECT department_id FROM departments WHERE department_code=?", (row["department_code"],)).fetchall()
            if not matches:
                errors.append(f"Row {index}: unknown department {row['department_code']}.")
            elif len(matches) > 1:
                errors.append(f"Row {index}: ambiguous department code {row['department_code']}; use a faculty-specific code.")

        elif dataset == "modules":
            if conn.execute("SELECT 1 FROM modules WHERE module_code=?", (row["module_code"],)).fetchone():
                errors.append(f"Row {index}: duplicate module {row['module_code']}.")
            matches = conn.execute("SELECT department_id FROM departments WHERE department_code=?", (row["department_code"],)).fetchall()
            if not matches:
                errors.append(f"Row {index}: unknown department {row['department_code']}.")
            elif len(matches) > 1:
                errors.append(f"Row {index}: ambiguous department code {row['department_code']}.")
            try:
                int(row["credits"])
            except ValueError:
                errors.append(f"Row {index}: credits must be an integer.")

        elif dataset == "student_groups":
            programme = conn.execute("SELECT programme_id FROM programmes WHERE programme_code=?", (row["programme_code"],)).fetchone()
            if programme is None:
                errors.append(f"Row {index}: unknown programme {row['programme_code']}.")
            try:
                year = int(row["year_level"])
                count = int(row["number_of_students"])
                if year <= 0 or count <= 0:
                    raise ValueError
            except ValueError:
                errors.append(f"Row {index}: year_level and number_of_students must be positive integers.")
            if programme is not None and conn.execute(
                "SELECT 1 FROM student_groups WHERE programme_id=? AND year_level=?",
                (programme["programme_id"], int(row["year_level"]) if row["year_level"].isdigit() else -1),
            ).fetchone():
                errors.append(f"Row {index}: duplicate programme/year group.")

        elif dataset == "offerings":
            module = conn.execute("SELECT module_id FROM modules WHERE module_code=?", (row["module_code"],)).fetchone()
            if module is None:
                errors.append(f"Row {index}: unknown module {row['module_code']}.")
            else:
                semester = conn.execute("SELECT semester_id FROM semesters WHERE status='ACTIVE' ORDER BY academic_year DESC, semester_number DESC LIMIT 1").fetchone()
                if semester and conn.execute("SELECT 1 FROM module_offerings WHERE module_id=? AND semester_id=?", (module["module_id"], semester["semester_id"])).fetchone():
                    errors.append(f"Row {index}: module {row['module_code']} already has an active-semester offering.")
            for token in [x.strip() for x in row["groups"].split(";") if x.strip()]:
                try:
                    pcode, year_text = token.split(":", 1)
                    found = conn.execute("""SELECT g.group_id FROM student_groups g JOIN programmes p ON p.programme_id=g.programme_id WHERE p.programme_code=? AND g.year_level=?""", (pcode, int(year_text))).fetchone()
                    if found is None:
                        errors.append(f"Row {index}: unknown student group {token}.")
                except Exception:
                    errors.append(f"Row {index}: group '{token}' must use PROGRAMME:YEAR, e.g. BCOMP:1.")
            for staff_number in [x.strip() for x in row["staff_numbers"].split(";") if x.strip()]:
                if conn.execute("SELECT 1 FROM staff WHERE staff_number=?", (staff_number,)).fetchone() is None:
                    errors.append(f"Row {index}: unknown staff number {staff_number}.")
            try:
                if int(row["expected_students"]) <= 0 or int(row["sessions_per_week"]) <= 0 or int(row["duration_slots"]) not in (1,2):
                    raise ValueError
                if int(row["can_split_duration"]) not in (0,1):
                    raise ValueError
            except ValueError:
                errors.append(f"Row {index}: invalid numeric offering fields.")

        normalized.append(row)

    if not rows:
        warnings.append("The uploaded file contains no data rows.")
    return {"dataset": dataset, "rows": normalized, "errors": errors, "warnings": warnings, "valid": not errors}


def commit_import(conn, preview):
    if not preview.get("valid"):
        raise ValueError("Cannot import invalid preview.")
    dataset = preview["dataset"]
    count = 0
    for row in preview["rows"]:
        if dataset == "rooms":
            building_id = conn.execute("SELECT building_id FROM buildings WHERE building_code=?", (row["building_code"],)).fetchone()[0]
            conn.execute(
                "INSERT INTO rooms(building_id, room_code, room_name, capacity, room_type) VALUES (?, ?, ?, ?, ?)",
                (building_id, row["room_code"], row["room_name"], int(row["capacity"]), row["room_type"]),
            )
        elif dataset == "staff":
            department_id = conn.execute("SELECT department_id FROM departments WHERE department_code=?", (row["department_code"],)).fetchone()[0]
            conn.execute(
                "INSERT INTO staff(department_id, staff_number, first_name, last_name, email) VALUES (?, ?, ?, ?, ?)",
                (department_id, row["staff_number"], row["first_name"], row["last_name"], row["email"]),
            )
        elif dataset == "modules":
            department_id = conn.execute("SELECT department_id FROM departments WHERE department_code=?", (row["department_code"],)).fetchone()[0]
            conn.execute(
                "INSERT INTO modules(department_id, module_code, module_name, credits) VALUES (?, ?, ?, ?)",
                (department_id, row["module_code"], row["module_name"], int(row["credits"])),
            )
        elif dataset == "student_groups":
            programme_id = conn.execute("SELECT programme_id FROM programmes WHERE programme_code=?", (row["programme_code"],)).fetchone()[0]
            conn.execute(
                "INSERT INTO student_groups(programme_id, year_level, number_of_students) VALUES (?, ?, ?)",
                (programme_id, int(row["year_level"]), int(row["number_of_students"])),
            )
        elif dataset == "offerings":
            semester_id = conn.execute("SELECT semester_id FROM semesters WHERE status='ACTIVE' ORDER BY academic_year DESC, semester_number DESC LIMIT 1").fetchone()[0]
            module_id = conn.execute("SELECT module_id FROM modules WHERE module_code=?", (row["module_code"],)).fetchone()[0]
            group_ids = []
            for token in [x.strip() for x in row["groups"].split(";") if x.strip()]:
                pcode, year_text = token.split(":", 1)
                group_ids.append(conn.execute("""SELECT g.group_id FROM student_groups g JOIN programmes p ON p.programme_id=g.programme_id WHERE p.programme_code=? AND g.year_level=?""", (pcode, int(year_text))).fetchone()[0])
            staff_ids = [conn.execute("SELECT staff_id FROM staff WHERE staff_number=?", (sn,)).fetchone()[0] for sn in [x.strip() for x in row["staff_numbers"].split(";") if x.strip()]]
            create_offering(
                conn, module_id, semester_id, int(row["expected_students"]), group_ids, staff_ids,
                row["session_type"], int(row["sessions_per_week"]), int(row["duration_slots"]),
                row["required_room_type"], bool(int(row["can_split_duration"])), []
            )
        count += 1
    conn.commit()
    return count
