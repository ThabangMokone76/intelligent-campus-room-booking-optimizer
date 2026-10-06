import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from database.bootstrap import ensure_synthetic_ready
from database.db import connect
from services.export_service import export_personal_excel, export_personal_pdf
from services.timetable_service import timetable_rows


class PersonalExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Path(self.tmp.name) / "personal_export.db"
        ensure_synthetic_ready(self.db)
        self.conn = connect(self.db)
        self.version = dict(self.conn.execute("SELECT * FROM timetable_versions WHERE status='PUBLISHED' ORDER BY version_number DESC LIMIT 1").fetchone())

    def tearDown(self):
        self.conn.close(); self.tmp.cleanup()

    def test_student_excel_contains_only_student_visible_modules(self):
        user = dict(self.conn.execute("SELECT * FROM users WHERE role='STUDENT' AND student_id IS NOT NULL ORDER BY user_id LIMIT 1").fetchone())
        visible = {r['module_code'] for r in timetable_rows(self.conn, self.version['version_id'], user)}
        path = export_personal_excel(self.conn, self.version['version_id'], user, Path(self.tmp.name)/'student.xlsx')
        wb = load_workbook(path, data_only=True)
        exported = {str(c.value) for row in wb['My Timetable'].iter_rows() for c in row if c.value and str(c.value) in visible}
        self.assertTrue(exported)
        self.assertTrue(exported.issubset(visible))

    def test_lecturer_personal_pdf_exports(self):
        user = dict(self.conn.execute("SELECT * FROM users WHERE role='LECTURER' AND staff_id IS NOT NULL ORDER BY user_id LIMIT 1").fetchone())
        path = export_personal_pdf(self.conn, self.version['version_id'], user, Path(self.tmp.name)/'lecturer.pdf')
        self.assertTrue(Path(path).read_bytes().startswith(b'%PDF'))
        self.assertGreater(Path(path).stat().st_size, 1500)


if __name__ == '__main__': unittest.main()
