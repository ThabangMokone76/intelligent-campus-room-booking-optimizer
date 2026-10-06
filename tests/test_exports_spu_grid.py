import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from database.bootstrap import ensure_synthetic_ready
from database.db import connect
from services.export_service import export_excel, export_pdf


class SPUGridExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Path(self.tmp.name) / "spu_export.db"
        ensure_synthetic_ready(self.db)
        self.conn = connect(self.db)
        row = self.conn.execute("SELECT version_id FROM timetable_versions WHERE status='PUBLISHED' ORDER BY version_number DESC LIMIT 1").fetchone()
        self.version_id = row["version_id"]

    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()

    def test_excel_uses_spu_day_venue_period_layout(self):
        path = export_excel(self.conn, self.version_id, Path(self.tmp.name) / "spu.xlsx")
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"])
        ws = wb["Monday"]
        values = [[cell.value for cell in row] for row in ws.iter_rows(min_row=1, max_row=8, min_col=1, max_col=12)]
        flat = [str(v) for row in values for v in row if v is not None]
        self.assertIn("VENUE", flat)
        self.assertIn("CAPACITY", flat)
        self.assertIn("P1", flat)
        self.assertIn("P10", flat)
        self.assertTrue(any("12:45-13:20" in v for v in flat))
        self.assertTrue(any(cell.alignment.wrap_text for row in ws.iter_rows() for cell in row if cell.value))

    def test_excel_and_pdf_tolerate_missing_faculty_colour(self):
        self.conn.execute("UPDATE faculties SET colour_hex=''  WHERE faculty_id=(SELECT MIN(faculty_id) FROM faculties)")
        self.conn.commit()
        xlsx = export_excel(self.conn, self.version_id, Path(self.tmp.name) / "neutral.xlsx")
        pdf = export_pdf(self.conn, self.version_id, Path(self.tmp.name) / "neutral.pdf")
        self.assertGreater(Path(xlsx).stat().st_size, 3000)
        self.assertTrue(Path(pdf).read_bytes().startswith(b"%PDF"))
        self.assertGreater(Path(pdf).stat().st_size, 5000)


    def test_excel_uses_canonical_faculty_colours(self):
        path = export_excel(self.conn, self.version_id, Path(self.tmp.name) / "colours.xlsx")
        wb = load_workbook(path)
        expected = {"059669", "7C3AED", "F6B221", "D97706"}
        observed = set()
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=6, min_col=3, max_col=12):
                for cell in row:
                    if cell.value and cell.fill and cell.fill.fill_type == "solid":
                        rgb = (cell.fill.fgColor.rgb or "").upper()
                        if len(rgb) == 8:
                            rgb = rgb[-6:]
                        observed.add(rgb)
        self.assertTrue(expected & observed)
        self.assertTrue(observed <= expected | {"FFFFFF", "E5E7EB"})


if __name__ == "__main__":
    unittest.main()
